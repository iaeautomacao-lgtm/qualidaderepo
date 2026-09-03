from __future__ import annotations

import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(r"C:\Users\gisele.oliveira\Downloads\QualiTalk_Formularios_Avaliacao.xlsx")
MD = Path(r"C:\Users\gisele.oliveira\Downloads\MAPEAMENTO_QualiTalk_Formularios.md")
OUT = ROOT / "database" / "cpanel" / "15-importa-formularios-qualitalk.sql"


CLIENTE_CAMPANHAS = {
    "Ânima": ["Chat", "Telefone ativo"],
    "B2B": ["B2B - Telefone"],
    "CEDAE": ["Telefone ativo Empresarial, Chat - Empresarial"],
    "Cobrança- Isaac": ["Isaac Ativo - Telefone"],
    "Cruzeiro do Sul": ["Chat", "Telefone Ativo"],
    "Educacional": ["Telefone"],
    "Empresarial - Cobrança": ["Chat - Empresarial", "Telefone ativo Empresarial"],
    "FIERGS": ["Ativo - Prospecção"],
    "FIESP": ["Telefone Ativo , Telefone Ativo, Chat"],
    "FIRJAN": [
        "Canais Online (E-mail, Chat e WhatsApp)",
        "E - Saúde Digital",
        "MONITORIAS IA",
        "Monitorias IA - Telefone Ativo",
        "Monitorias IA - Telefone Receptivo",
        "Odontologia e Massagem relaxante (Offline)",
        "Telefone Ativo",
        "Telefone Receptivo",
    ],
    "Grupo Avenida": ["Telefone Grupo Avenida"],
    "Receptivo": ["Chat", "Telefone Receptivo"],
    "teste 1": ["teste 1 campanha"],
    "Vero": [
        "Ativo 20 a 44",
        "Ativo 45 a 75",
        "Ativo 76 a 120",
        "B2B",
        "Churn",
        "Pré churn",
        "Pré Churn",
        "Vero Ativo",
    ],
    "Yduqs": ["Telefone Ativo"],
}


def normalizar(texto: str) -> str:
    return (
        unicodedata.normalize("NFD", texto or "")
        .encode("ascii", "ignore")
        .decode("ascii")
        .casefold()
        .strip()
    )


def slug(texto: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", normalizar(texto)).strip("-")
    return base[:110] or "cliente"


def sql(texto) -> str:
    if texto is None:
        return "NULL"
    return "'" + str(texto).replace("\\", "\\\\").replace("'", "''") + "'"


def inferir_cliente(nome: str) -> str:
    n = normalizar(nome)
    if "firjan" in n:
        return "FIRJAN"
    if "cedae" in n:
        return "CEDAE"
    if "fiesp" in n:
        return "FIESP"
    if "fiergs" in n:
        return "FIERGS"
    if nome == "Vero B2B":
        return "B2B"
    if "vero" in n:
        return "Vero"
    if "grupo avenida" in n:
        return "Grupo Avenida"
    if "isaac" in n:
        return "Cobrança- Isaac"
    if "cruzeiro" in n:
        return "Cruzeiro do Sul"
    if "anima" in n:
        return "Ânima"
    if "yduqs" in n:
        return "Yduqs"
    if "educacional | receptivo" in n:
        return "Receptivo"
    if "avaliacao cobranca educacional" in n:
        return "Educacional"
    if "receptivo - educacional - empresarial" in n:
        return "Receptivo"
    if "empresarial - cobranca" in n:
        return "Empresarial - Cobrança"
    raise ValueError(f"Cliente não mapeado para formulário: {nome}")


def canal(campanha: str) -> str:
    n = normalizar(campanha)
    if "chat" in n or "whatsapp" in n or "online" in n:
        return "chat"
    if "email" in n:
        return "email"
    if "offline" in n:
        return "offline"
    if "telefone" in n or "ativo" in n or "receptivo" in n:
        return "telefone"
    return "outro"


def separar_campanhas(valor: str) -> list[str]:
    if not valor or valor == "—":
        return []
    partes = [p.strip() for p in valor.split(",") if p.strip()]
    vistas = set()
    saida = []
    for parte in partes:
        chave = normalizar(parte)
        if chave not in vistas:
            vistas.add(chave)
            saida.append(parte)
    return saida


def campanhas_do_formulario(cliente: str, nome: str, bruto: str) -> list[str]:
    permitidas = CLIENTE_CAMPANHAS[cliente]
    permitidas_norm = {normalizar(item): item for item in permitidas}
    n = normalizar(nome)

    if cliente in {"CEDAE", "FIESP"}:
        return permitidas
    if cliente == "Vero" and "ia - churn" in n:
        return ["Churn"]
    if cliente == "Vero" and "pre churn" in n and "ia" in n:
        return ["Pré Churn"]
    if cliente == "Vero" and "sp - pre churn" in n:
        return ["Pré Churn"]
    if cliente == "Vero" and "sp - ativo" in n:
        return ["Vero Ativo"]

    escolhidas = []
    for campanha in separar_campanhas(bruto):
        encontrada = permitidas_norm.get(normalizar(campanha))
        if encontrada and encontrada not in escolhidas:
            escolhidas.append(encontrada)
    return escolhidas or permitidas[:1]


def categoria(valor: str, nome: str) -> str:
    texto = normalizar(f"{valor} {nome}")
    return "diagnostico" if "diagnostico" in texto or " ia" in f" {texto} " else "padrao"


def status(valor: str) -> str:
    return "ativo" if valor == "active" else "inativo"


def descricoes_markdown() -> dict[str, str]:
    if not MD.exists():
        return {}
    texto = MD.read_text(encoding="utf-8")
    descricoes = {}
    partes = re.split(r"\n\*\*▸ ", texto)
    for parte in partes[1:]:
        nome = parte.split("**", 1)[0].strip()
        match = re.search(r"\n> ([^\n]+)", parte)
        if match:
            descricoes[nome] = match.group(1).strip()
    return descricoes


def secao_nome_descricao(valor: str) -> tuple[str, str | None]:
    texto = str(valor or "").strip()
    if "\n_" not in texto:
        return texto, None
    nome, descricao = texto.split("\n_", 1)
    return nome.strip(), descricao.strip().strip("_") or None


def carregar():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws_forms = wb["Formulários"]
    ws_criterios = wb["Critérios"]
    descricoes = descricoes_markdown()

    forms = []
    for row in ws_forms.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        nome = str(row[1]).strip()
        cliente = inferir_cliente(nome)
        forms.append(
            {
                "nome": nome,
                "source_id": str(row[2]).strip(),
                "status": status(str(row[3]).strip()),
                "categoria": categoria(str(row[4] or ""), nome),
                "campanhas": campanhas_do_formulario(cliente, nome, str(row[5] or "")),
                "cliente": cliente,
                "bloco": str(row[10]).strip(),
                "descricao": descricoes.get(nome),
            }
        )

    criterios = defaultdict(list)
    for pos, row in enumerate(ws_criterios.iter_rows(min_row=2, values_only=True), start=1):
        if not row[0] or not row[2] or not row[3]:
            continue
        secao, secao_descricao = secao_nome_descricao(row[2])
        criterio = str(row[3]).strip()
        descricao = str(row[6] or criterio).strip()
        eliminatoria = str(row[5] or "").strip().upper() == "SIM"
        peso = None if eliminatoria else float(row[4] or 0)
        criterios[str(row[0]).strip()].append(
            {
                "secao": secao,
                "secao_descricao": secao_descricao,
                "criterio": criterio[:200],
                "descricao": descricao,
                "peso": peso,
                "eliminatoria": eliminatoria,
                "ordem": pos,
            }
        )
    return forms, criterios


def gerar_sql(forms, criterios) -> str:
    linhas = [
        "-- Importação controlada dos formulários QualiTalk",
        "-- Fonte: QualiTalk_Formularios_Avaliacao.xlsx / MAPEAMENTO_QualiTalk_Formularios.md",
        "-- Escopo: cria/atualiza formulários por nome + cliente; não apaga formulários fora desta lista.",
        "-- Se um formulário existente já tem avaliações, uma nova versão é criada para preservar as respostas antigas.",
        "",
        "SET NAMES utf8mb4;",
        "START TRANSACTION;",
        "",
        "SET @schema_name := DATABASE();",
        "SET @tem_descricao := (SELECT COUNT(*) FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=@schema_name AND TABLE_NAME='formularios' AND COLUMN_NAME='descricao');",
        "SET @ddl := IF(@tem_descricao=0, 'ALTER TABLE formularios ADD COLUMN descricao TEXT NULL AFTER nome', 'SELECT 1');",
        "PREPARE stmt FROM @ddl; EXECUTE stmt; DEALLOCATE PREPARE stmt;",
        "SET @tem_tipo_calculo := (SELECT COUNT(*) FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=@schema_name AND TABLE_NAME='formularios' AND COLUMN_NAME='tipo_calculo');",
        "SET @ddl := IF(@tem_tipo_calculo=0, \"ALTER TABLE formularios ADD COLUMN tipo_calculo ENUM(''sessao'',''criterio'') NOT NULL DEFAULT ''criterio'' AFTER categoria\", 'SELECT 1');",
        "PREPARE stmt FROM @ddl; EXECUTE stmt; DEALLOCATE PREPARE stmt;",
        "",
    ]

    for form in forms:
        linhas.extend(
            [
                f"-- {form['cliente']} / {form['nome']} / QualiTalk {form['source_id']}",
                f"SET @cliente_nome := {sql(form['cliente'])};",
                f"SET @cliente_slug := {sql(slug(form['cliente']))};",
                "INSERT INTO clientes (slug, nome, ativo) VALUES (@cliente_slug, @cliente_nome, 1) ON DUPLICATE KEY UPDATE nome=VALUES(nome), ativo=1;",
                "SET @cliente_id := (SELECT id FROM clientes WHERE slug=@cliente_slug OR nome=@cliente_nome ORDER BY id LIMIT 1);",
                f"SET @form_nome := {sql(form['nome'])};",
                f"SET @form_descricao := {sql(((form['descricao'] or '') + ' [QualiTalk ID: ' + form['source_id'] + ']').strip())};",
                f"SET @form_categoria := {sql(form['categoria'])};",
                f"SET @form_status := {sql(form['status'])};",
                "SET @form_existente_id := (SELECT id FROM formularios WHERE cliente_id=@cliente_id AND nome=@form_nome ORDER BY versao DESC, id DESC LIMIT 1);",
                "SET @form_tem_avaliacoes := IFNULL((SELECT COUNT(*) FROM avaliacoes WHERE formulario_id=@form_existente_id), 0);",
                "SET @form_versao := IF(@form_existente_id IS NULL, 1, IF(@form_tem_avaliacoes=0, (SELECT versao FROM formularios WHERE id=@form_existente_id), (SELECT COALESCE(MAX(versao),0)+1 FROM formularios WHERE cliente_id=@cliente_id AND nome=@form_nome)));",
                "UPDATE formularios SET descricao=@form_descricao, tipo_calculo='sessao', categoria=@form_categoria, status=@form_status WHERE id=@form_existente_id AND @form_tem_avaliacoes=0;",
                "INSERT INTO formularios (cliente_id, nome, descricao, categoria, tipo_calculo, status, versao)",
                "SELECT @cliente_id, @form_nome, @form_descricao, @form_categoria, 'sessao', @form_status, @form_versao",
                "WHERE @form_existente_id IS NULL OR @form_tem_avaliacoes > 0;",
                "SET @formulario_id := IF(@form_existente_id IS NULL OR @form_tem_avaliacoes > 0, LAST_INSERT_ID(), @form_existente_id);",
                "UPDATE formularios SET status='inativo' WHERE cliente_id=@cliente_id AND nome=@form_nome AND id<>@formulario_id AND @form_tem_avaliacoes > 0;",
                "DELETE FROM formulario_campanhas WHERE formulario_id=@formulario_id;",
                "DELETE FROM formulario_secoes WHERE formulario_id=@formulario_id;",
            ]
        )

        for campanha in form["campanhas"]:
            linhas.extend(
                [
                    f"SET @campanha_nome := {sql(campanha)};",
                    f"INSERT INTO campanhas (cliente_id, nome, canal, ativa) VALUES (@cliente_id, @campanha_nome, {sql(canal(campanha))}, 1) ON DUPLICATE KEY UPDATE canal=VALUES(canal), ativa=1;",
                    "SET @campanha_id := (SELECT id FROM campanhas WHERE cliente_id=@cliente_id AND nome=@campanha_nome LIMIT 1);",
                    "INSERT IGNORE INTO formulario_campanhas (formulario_id, campanha_id) VALUES (@formulario_id, @campanha_id);",
                ]
            )

        secoes = []
        vistos = set()
        for criterio in criterios[form["bloco"]]:
            chave = (criterio["secao"], criterio["secao_descricao"])
            if chave not in vistos:
                vistos.add(chave)
                secoes.append(chave)

        for secao_pos, (secao, secao_descricao) in enumerate(secoes, start=1):
            linhas.extend(
                [
                    f"INSERT INTO formulario_secoes (formulario_id, nome, descricao, posicao) VALUES (@formulario_id, {sql(secao)}, {sql(secao_descricao)}, {secao_pos});",
                    "SET @secao_id := LAST_INSERT_ID();",
                ]
            )
            criterio_pos = 1
            for criterio in criterios[form["bloco"]]:
                if (criterio["secao"], criterio["secao_descricao"]) != (secao, secao_descricao):
                    continue
                peso = "NULL" if criterio["peso"] is None else f"{criterio['peso']:.2f}"
                elim = "1" if criterio["eliminatoria"] else "0"
                linhas.append(
                    "INSERT INTO formulario_criterios (secao_id, nome, enunciado, peso_pts, eliminatoria, posicao) "
                    f"VALUES (@secao_id, {sql(criterio['criterio'])}, {sql(criterio['descricao'])}, {peso}, {elim}, {criterio_pos});"
                )
                criterio_pos += 1

        linhas.append("")

    linhas.extend(
        [
            "COMMIT;",
            "",
            "SELECT f.id, cl.nome AS cliente, f.nome, f.versao, f.status, COUNT(DISTINCT cr.id) AS criterios",
            "  FROM formularios f",
            "  JOIN clientes cl ON cl.id = f.cliente_id",
            "  LEFT JOIN formulario_secoes s ON s.formulario_id = f.id",
            "  LEFT JOIN formulario_criterios cr ON cr.secao_id = s.id",
            " WHERE f.nome IN (",
            "   " + ",\n   ".join(sql(form["nome"]) for form in forms),
            " )",
            " GROUP BY f.id, cl.nome, f.nome, f.versao, f.status",
            " ORDER BY cl.nome, f.nome, f.versao;",
            "",
        ]
    )
    return "\n".join(linhas)


def main() -> int:
    if not XLSX.exists():
        print(f"Planilha não encontrada: {XLSX}", file=sys.stderr)
        return 1
    forms, criterios = carregar()
    total = sum(len(criterios[form["bloco"]]) for form in forms)
    OUT.write_text(gerar_sql(forms, criterios), encoding="utf-8")
    print(f"Gerado {OUT}")
    print(f"Formulários: {len(forms)}")
    print(f"Critérios gerados por formulário: {total}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
